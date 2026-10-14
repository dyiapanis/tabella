"""Regression test: section type 'text' renders as a paragraph (bug fix 2026-10-15).

Previously, `{"type": "text", ...}` sections fell through the builder's elif
chain unchecked (docx/odt) → silent EMPTY document. PDF handled it; the
document builders did not. Fix: normalize "text" → "paragraph" at the section
loop in both builders.

Run: python3 tests/test_text_section_alias.py  (needs docs venv python for
python-docx / odfpy imports)
"""
import os
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from tabella.create import create_document  # noqa: E402


def test_docx_text_alias():
    from docx import Document

    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "a.docx")
        r = create_document(
            format="docx",
            content={"sections": [{"type": "text", "text": "ALIAS-PAYLOAD"}]},
            output_path=p,
        )
        assert r["success"], r
        texts = [pp.text for pp in Document(p).paragraphs if pp.text.strip()]
        assert texts == ["ALIAS-PAYLOAD"], texts


def test_odt_text_alias():
    from odf.opendocument import load
    from odf import text as odftext, teletype

    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "a.odt")
        r = create_document(
            format="odt",
            content={"sections": [{"type": "text", "text": "ALIAS-PAYLOAD"}]},
            output_path=p,
        )
        assert r["success"], r
        paras = load(p).getElementsByType(odftext.P)
        assert any("ALIAS-PAYLOAD" in teletype.extractText(x) for x in paras)


def test_docx_known_types_regress_none():
    from docx import Document

    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "b.docx")
        create_document(
            format="docx",
            content={"sections": [
                {"type": "heading", "text": "H1", "level": 1},
                {"type": "paragraph", "text": "P1"},
                {"type": "table", "headers": ["col"], "rows": [["v"]]},
            ]},
            output_path=p,
        )
        doc = Document(p)
        body = "\n".join(pp.text for pp in doc.paragraphs)
        cells = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
        assert "H1" in body and "P1" in body
        assert "col" in cells and "v" in cells


if __name__ == "__main__":
    test_docx_text_alias()
    test_odt_text_alias()
    test_docx_known_types_regress_none()
    print("3/3 tests PASS")